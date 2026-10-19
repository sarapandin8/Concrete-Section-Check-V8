"""SLAB.IMPORT55: input-only CSiBridge shell package, separate from girder loads.

Read complete workbooks without requiring renamed files or edited worksheets.
This module validates import correspondence, not slab resistance or design.
"""
from __future__ import annotations

from dataclasses import dataclass, field
from io import BytesIO
import hashlib
import re
from typing import Any

import numpy as np
import openpyxl
import pandas as pd

VERSION = "SLAB.IMPORT55"
BANK_KEY = "slab_shell55_package"
PENDING_KEY = "slab_shell55_selected_signature"
FORCES = ("F11", "F22", "F12", "M11", "M22", "M12", "V13", "V23")


@dataclass(frozen=True)
class TableSpec:
    key: str
    title: str
    columns: tuple[str, ...]
    group: int
    conditional: bool = False


SPECS = (
    TableSpec("forces", "Element Forces - Area Shells",
        ("Area", "AreaElem", "ShellType", "Joint", "OutputCase", "CaseType", "StepType", *FORCES), 1),
    TableSpec("elements", "Objects And Elements - Areas",
        ("AreaElem", "AreaObject", "ElemJt1", "ElemJt2", "ElemJt3", "ElemJt4"), 2),
    TableSpec("nodes", "Objects And Elements - Joints", ("JointElem", "GlobalX", "GlobalY", "GlobalZ"), 2),
    TableSpec("flags", "Area Bridge Object Flags", ("Area", "BridgeObj", "BOSpan", "CompType", "StartSta", "EndSta"), 3),
    TableSpec("thickness", "Area Overwrites - Thickness", ("Area", "ThickOver", "Thick1", "Thick2", "Thick3", "Thick4"), 3),
    TableSpec("sections", "Area Section Assignments", ("Area", "Section", "MatProp"), 3),
    TableSpec("axes", "Area Local Axes Assignments 1 - Typical", ("Area", "Angle", "AdvanceAxes"), 3),
    TableSpec("advanced", "Area Local Axes Assignments 2 - Advanced",
        ("Area", "LocalPlane", "PlOption1", "PlCoordSys", "CoordDir1", "CoordDir2", "PlVecJt1", "PlVecJt2", "PlVecX", "PlVecY", "PlVecZ"), 3, True),
    TableSpec("parents", "Connectivity - Area", ("Area", "NumJoints", "Joint1", "Joint2", "Joint3", "Joint4"), 4),
    TableSpec("joints", "Joint Coordinates", ("Joint", "GlobalX", "GlobalY", "GlobalZ"), 4),
)

FILE_GUIDE = (
    {"File / export set": "Element Forces - Area Shells.xlsx", "Required tables": "Element Forces - Area Shells",
     "Use": "ULS: F11, F22, F12, M11, M22, M12, V13, V23 · all supplied cases and Max/Min rows"},
    {"File / export set": "TABLE Objects And Elements - Areas(1).xlsx", "Required tables": "Objects And Elements - Areas; Objects And Elements - Joints",
     "Use": "Analysis mesh connectivity and coordinates, including generated ~ joints"},
    {"File / export set": "Area Assigment.xlsx", "Required tables": "Area Bridge Object Flags; Area Overwrites - Thickness; Area Section Assignments; Area Local Axes Assignments 1 - Typical; Area Local Axes Assignments 2 - Advanced*",
     "Use": "Top Slab selection, bridge stations, corner thicknesses, section assignment and shell local axes"},
    {"File / export set": "Connectivity Data.xlsx", "Required tables": "Connectivity - Area; Joint Coordinates",
     "Use": "Original Area corner order, original coordinates and axis reference joints"},
)


def _token(value: Any) -> str:
    return re.sub(r"[^a-z0-9]", "", str(value or "").casefold())


def identifier(value: Any) -> str:
    if value is None or (isinstance(value, float) and np.isnan(value)):
        return ""
    if isinstance(value, (int, float)) and not isinstance(value, bool) and float(value).is_integer():
        return str(int(value))
    return str(value).strip()


def _unit(value: Any) -> str:
    return str(value or "").casefold().replace(" ", "").replace("·", "-").replace("²", "2")


@dataclass
class ShellPackage:
    tables: dict[str, pd.DataFrame] = field(default_factory=dict)
    units: dict[str, dict[str, str]] = field(default_factory=dict)
    sources: list[dict[str, Any]] = field(default_factory=list)
    errors: list[str] = field(default_factory=list)
    notes: list[str] = field(default_factory=list)
    signature: tuple[str, ...] = ()
    summary: dict[str, Any] = field(default_factory=dict)

    def required(self) -> list[TableSpec]:
        axes = self.tables.get("axes", pd.DataFrame())
        flags = self.tables.get("flags", pd.DataFrame())
        if not flags.empty and not axes.empty:
            slab_ids = flags.loc[flags["CompType"].astype(str).str.strip().str.casefold().eq("top slab"), "Area"].map(identifier)
            axes = axes.loc[axes["Area"].map(identifier).isin(slab_ids)]
        advanced = not axes.empty and axes["AdvanceAxes"].astype(str).str.casefold().eq("yes").any()
        return [s for s in SPECS if not s.conditional or advanced]

    @property
    def missing(self) -> list[TableSpec]:
        return [s for s in self.required() if self.tables.get(s.key, pd.DataFrame()).empty]

    @property
    def ready(self) -> bool:
        return not self.missing and not self.errors and bool(self.summary.get("slab_force_rows"))


def upload_signature(files) -> tuple[str, ...]:
    return tuple(sorted(hashlib.sha256(f.getvalue()).hexdigest() for f in files))


def _read_book(payload: bytes, name: str, package: ShellPackage) -> None:
    book = openpyxl.load_workbook(BytesIO(payload), read_only=True, data_only=True)
    try:
        for sheet in book:
            rows = sheet.iter_rows(values_only=True)
            spec = None
            columns: list[str] = []
            first_title = sheet.title
            for index, row in enumerate(rows):
                texts = [str(x or "").strip() for x in row]
                if texts and texts[0].upper().startswith("TABLE:"):
                    first_title = texts[0].split(":", 1)[1].strip()
                candidates = [s for s in SPECS if set(s.columns).issubset(texts)]
                if candidates:
                    spec = candidates[0]
                    # Ignore trailing empty columns while preserving all exported fields.
                    last = max(i for i, c in enumerate(texts) if c)
                    columns = texts[:last + 1]
                    if not all(columns) or len(set(columns)) != len(columns):
                        raise ValueError(f"{sheet.title}: blank or duplicate column headers")
                    break
                if index >= 9:
                    break
            if spec is None:
                # Extra CSI tables are permitted. A known but malformed table is actionable.
                aliases = {_token(sheet.title), _token(first_title)}
                known = next((s for s in SPECS if _token(s.title) in aliases or
                    _token(s.title.replace(" Assignments", "")) in aliases), None)
                if known:
                    package.errors.append(f"{name} / {sheet.title}: required columns: {', '.join(known.columns)}")
                package.sources.append({"File": name, "Table": sheet.title, "Use": "Not used", "Rows": None})
                continue
            unit_row = next(rows, ())
            units = {c: str(v or "").strip() for c, v in zip(columns, unit_row)}
            # CSI native exports contain a typed units row directly after headers.
            if not unit_row or str(unit_row[0]).strip().casefold() != "text":
                raise ValueError(f"{sheet.title}: expected the original CSI units row below the headers; export the complete table")
            records = [row[:len(columns)] for row in rows if any(v is not None for v in row)
                and not str(row[0] or "").upper().startswith("END TABLE")]
            frame = pd.DataFrame(records, columns=columns)
            if spec.key in package.tables:
                previous = package.tables[spec.key]
                if units == package.units[spec.key] and frame.equals(previous):
                    package.notes.append(f"Identical duplicate table ignored: {name} / {sheet.title}")
                else:
                    package.errors.append(f"Conflicting copies of {spec.title}; select one consistent model export")
            else:
                package.tables[spec.key], package.units[spec.key] = frame, units
            package.sources.append({"File": name, "Table": sheet.title, "Use": spec.title, "Rows": len(frame)})
    finally:
        book.close()


def parse_files(files) -> ShellPackage:
    """Input audit only; never apply envelopes, convert moments or run a solver."""
    package = ShellPackage(signature=upload_signature(files))
    seen: set[str] = set()
    for upload in files:
        payload = upload.getvalue()
        digest = hashlib.sha256(payload).hexdigest()
        if digest in seen:
            package.notes.append(f"Identical file ignored: {upload.name}")
            continue
        seen.add(digest)
        try:
            _read_book(payload, str(upload.name), package)
        except Exception as exc:
            # A bad workbook must produce a visible blocked import, never partial readiness.
            package.errors.append(f"{upload.name}: {type(exc).__name__}: {exc}")
    _validate(package)
    return package


def _validate(package: ShellPackage) -> None:
    tables = package.tables
    id_columns = ("Area", "AreaElem", "AreaObject", "Joint", "JointElem", "ElemJt1", "ElemJt2", "ElemJt3", "ElemJt4",
        "Joint1", "Joint2", "Joint3", "Joint4", "PlVecJt1", "PlVecJt2")
    for frame in tables.values():
        for col in set(frame.columns).intersection(id_columns):
            frame[col] = frame[col].map(identifier)
    for key, id_col in (("elements", "AreaElem"), ("nodes", "JointElem"), ("flags", "Area"), ("thickness", "Area"),
                        ("sections", "Area"), ("axes", "Area"), ("advanced", "Area"), ("parents", "Area"), ("joints", "Joint")):
        if key in tables and (tables[key][id_col].eq("").any() or tables[key][id_col].duplicated().any()):
            package.errors.append(f"{key}: blank or duplicate {id_col}")
    for key, columns in (("forces", FORCES), ("nodes", ("GlobalX", "GlobalY", "GlobalZ")),
                         ("joints", ("GlobalX", "GlobalY", "GlobalZ")), ("axes", ("Angle",)),
                         ("thickness", ("Thick1", "Thick2", "Thick3", "Thick4")), ("flags", ("StartSta", "EndSta"))):
        if key not in tables:
            continue
        for column in columns:
            values = pd.to_numeric(tables[key][column], errors="coerce")
            if not np.isfinite(values.to_numpy(dtype=float)).all() or tables[key][column].map(lambda v: isinstance(v, bool)).any():
                package.errors.append(f"{key}: {column} contains blank or non-finite numbers")
            tables[key][column] = values
    # Explicitly scoped to the units in the supplied complete CSI exports.
    checks = [("forces", c, {"kn/m"} if c[0] != "M" else {"kn-m/m", "knm/m"}) for c in FORCES]
    checks += [(k, c, {"m"}) for k in ("nodes", "joints") for c in ("GlobalX", "GlobalY", "GlobalZ")]
    checks += [("thickness", f"Thick{i}", {"m"}) for i in range(1, 5)]
    checks += [("axes", "Angle", {"degrees", "deg"})]
    checks += [("flags", c, {"m"}) for c in ("StartSta", "EndSta")]
    for key, column, allowed in checks:
        if key in tables and _unit(package.units[key].get(column)) not in allowed:
            package.errors.append(f"{key}: unsupported unit for {column}; export from CSI in kN, m (angles: degrees)")
    if "forces" in tables:
        forces = tables["forces"]
        for col in ("OutputCase", "CaseType", "StepType"):
            forces[col] = forces[col].fillna("").astype(str).str.strip()
        if forces[["Area", "AreaElem", "Joint", "OutputCase"]].eq("").any().any():
            package.errors.append("forces: blank Area / AreaElem / Joint / OutputCase")
        keys = ["Area", "AreaElem", "Joint", "OutputCase", "StepType"]
        if "StepNum" in forces:
            keys.append("StepNum")
        if forces.duplicated(keys).any():
            package.errors.append("forces: duplicate area / element / corner / case / step records")
    if package.missing or package.errors:
        return
    forces, elements, nodes, flags = (tables[k] for k in ("forces", "elements", "nodes", "flags"))
    slab_ids = set(flags.loc[flags["CompType"].astype(str).str.strip().str.casefold().eq("top slab"), "Area"])
    if not slab_ids:
        package.errors.append("Area Bridge Object Flags: no Top Slab objects found")
        return
    for key in ("thickness", "sections", "axes", "parents"):
        absent = slab_ids - set(tables[key]["Area"])
        if absent:
            package.errors.append(f"{key}: {len(absent)} Top Slab objects missing (e.g. {', '.join(sorted(absent)[:4])})")
    slab_axes = tables["axes"].loc[tables["axes"]["Area"].isin(slab_ids)]
    needs_advanced = set(slab_axes.loc[slab_axes["AdvanceAxes"].astype(str).str.casefold().eq("yes"), "Area"])
    if needs_advanced - set(tables.get("advanced", pd.DataFrame(columns=["Area"]))["Area"]):
        package.errors.append("advanced: assignments missing for Top Slab objects with AdvanceAxes = Yes")
    joined = forces.merge(elements, on="AreaElem", how="left", validate="many_to_one")
    if joined["AreaObject"].isna().any():
        package.errors.append("forces: AreaElem not found in Objects And Elements - Areas")
    if not joined["Area"].eq(joined["AreaObject"]).all():
        package.errors.append("forces: parent Area differs from mesh AreaObject")
    corners = [f"ElemJt{i}" for i in range(1, 5)]
    matches = joined[corners].eq(joined["Joint"], axis=0).any(axis=1)
    if not matches.all():
        package.errors.append("forces: Joint is not a corner of its AreaElem")
    mesh_joints = set(elements[corners].to_numpy().ravel()) - {""}
    if mesh_joints - set(nodes["JointElem"]):
        package.errors.append("nodes: coordinates missing for analysis mesh joints, including generated ~ joints")
    if set(elements["AreaObject"]) - set(flags["Area"]):
        package.errors.append("flags: component classification missing for mesh Area objects")
    slab_parents = tables["parents"].loc[tables["parents"]["Area"].isin(slab_ids)]
    original_nodes = set(tables["joints"]["Joint"]) | set(nodes["JointElem"])
    parent_joints = set(slab_parents[[f"Joint{i}" for i in range(1, 5)]].to_numpy().ravel()) - {""}
    if parent_joints - original_nodes:
        package.errors.append("joints: original Area corner coordinates missing")
    if not pd.to_numeric(slab_parents["NumJoints"], errors="coerce").eq(4).all():
        package.errors.append("parents: this import route requires four-corner Top Slab Areas")
    if "advanced" in tables:
        adv = tables["advanced"].loc[tables["advanced"]["Area"].isin(needs_advanced)]
        two_joints = adv["PlOption1"].astype(str).str.casefold().eq("two joints")
        refs = set(adv.loc[two_joints, ["PlVecJt1", "PlVecJt2"]].to_numpy().ravel())
        if "" in refs or refs - original_nodes:
            package.errors.append("advanced: Two Joints reference coordinates missing")
    thickness = tables["thickness"].loc[tables["thickness"]["Area"].isin(slab_ids)]
    if not thickness["ThickOver"].astype(str).str.casefold().eq("object").all():
        package.errors.append("thickness: this export route requires explicit object-corner thicknesses")
    values = thickness[[f"Thick{i}" for i in range(1, 5)]].to_numpy(dtype=float)
    if not (values > 0).all():
        package.errors.append("thickness: non-positive Top Slab corner thickness")
    slab = forces.loc[forces["Area"].isin(slab_ids)]
    if slab.empty:
        package.errors.append("forces: no records match Top Slab objects")
    package.summary = {
        "force_rows": len(forces), "slab_force_rows": len(slab), "excluded_force_rows": len(forces) - len(slab),
        "slab_objects": len(slab_ids), "mesh_elements": len(elements),
        "slab_elements": int(elements["AreaObject"].isin(slab_ids).sum()), "analysis_joints": len(nodes),
        "thickness_min_mm": float(values.min() * 1000) if values.size else None,
        "thickness_max_mm": float(values.max() * 1000) if values.size else None,
        "cases": sorted(slab["OutputCase"].unique().tolist()),
        "case_rows": slab.groupby(["OutputCase", "StepType"], dropna=False).size().reset_index(name="Rows"),
    }


def current_package(state) -> ShellPackage | None:
    package = state.get(BANK_KEY)
    if not isinstance(package, ShellPackage) or not package.ready:
        return None
    if state.get(PENDING_KEY, package.signature) != package.signature:
        return None
    return package
