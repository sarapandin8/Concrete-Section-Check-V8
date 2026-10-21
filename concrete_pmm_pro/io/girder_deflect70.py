"""Input contract for CSI Envelope SLS Girder Deflection workbooks.

One worksheet identifies one girder. Only Girder Distance, ItemType and
WebTop Vert enter the input bank. Max/Min pairs establish the envelope shape;
the SLS load-regime declaration is supplied by this dedicated import route.
The illustrated CSI export does not contain a load-combination definition.
"""
from __future__ import annotations

from copy import deepcopy
from hashlib import sha256
from io import BytesIO
import math
from pathlib import Path
import re
from zipfile import BadZipFile

from openpyxl import load_workbook

BANK_KEY = 'girder_deflect70_input'
ERROR_KEY = 'girder_deflect70_upload_error'
SCHEMA = 'CSI_SLS_GIRDER_ENVELOPE_1'
REQUIRED = ('Girder Distance', 'ItemType', 'WebTop Vert')
UNIT_TO_M = {'m': 1., 'mm': .001, 'cm': .01, 'ft': .3048, 'in': .0254}


def _text(value):
    return ' '.join(str(value or '').replace('\xa0', ' ').split())


def _number(value, location):
    if isinstance(value, bool) or value is None:
        raise ValueError(location + ': a finite numeric value is required')
    try:
        number = float(value)
    except (ValueError, TypeError):
        raise ValueError(location + ': a finite numeric value is required') from None
    if not math.isfinite(number):
        raise ValueError(location + ': a finite numeric value is required')
    return number


def _unit(value, location):
    unit = _text(value).casefold()
    if unit not in UNIT_TO_M:
        raise ValueError(location + ': length unit must be m, mm, cm, ft or in')
    return unit


def _validate_sheet(sheet):
    if not isinstance(sheet, dict):
        raise ValueError('Each girder input must contain a worksheet record')
    name = _text(sheet.get('girder'))
    if not name or re.search(r'entire\s+(bridge|section)|all\s+girders|bridge\s+section', name, re.I):
        raise ValueError('Use an individual girder worksheet, not the entire bridge section')
    distance_unit = _unit(sheet.get('distance_unit'), name + ' / Girder Distance')
    vertical_unit = _unit(sheet.get('vertical_unit'), name + ' / WebTop Vert')
    records = sheet.get('rows')
    if not isinstance(records, list) or not records:
        raise ValueError(name + ': Max/Min envelope rows are required')
    stations = {}
    rows = []
    for record in records:
        if not isinstance(record, dict):
            raise ValueError(name + ': each source row must be a record')
        row = _number(record.get('source_row'), name + ' / source row')
        if row < 1 or row != int(row):
            raise ValueError(name + ': source row must be a positive integer')
        loc = f'{name}, row {int(row)}'
        branch = _text(record.get('ItemType')).title()
        if branch not in {'Max', 'Min'}:
            raise ValueError(loc + ': only Envelope Max/Min is accepted; single-case/step results are not supported')
        x = _number(record.get('Girder Distance'), loc + ' / Girder Distance')
        vertical = _number(record.get('WebTop Vert'), loc + ' / WebTop Vert')
        if x < 0:
            raise ValueError(loc + ': Girder Distance cannot be negative')
        pair = stations.setdefault(x, {})
        if branch in pair:
            raise ValueError(loc + ': duplicate station / ItemType; retain one exported result for each Max/Min branch')
        pair[branch] = vertical
        rows.append({'source_row': int(row), 'Girder Distance': x,
            'ItemType': branch, 'WebTop Vert': vertical})
    if len(stations) < 2:
        raise ValueError(name + ': at least two stations are required')
    for x, pair in stations.items():
        if set(pair) != {'Max', 'Min'}:
            raise ValueError(f'{name}, x={x:g}: both Envelope Max and Min are required')
        if pair['Max'] < pair['Min'] - 1e-12:
            raise ValueError(f'{name}, x={x:g}: Envelope Max is below Min; check the export')
    return {'girder': name, 'distance_unit': distance_unit, 'vertical_unit': vertical_unit, 'rows': rows}


def validate_bank(bank):
    if not isinstance(bank, dict) or bank.get('schema') != SCHEMA:
        raise ValueError('Unsupported SLS girder deflection input schema')
    sheets = bank.get('girders')
    if not isinstance(sheets, list) or not sheets:
        raise ValueError('No individual-girder SLS envelope was found')
    sheets = [_validate_sheet(sheet) for sheet in sheets]
    names = [sheet['girder'].casefold() for sheet in sheets]
    if len(names) != len(set(names)):
        raise ValueError('Girder worksheet names must be unique')
    digest = str(bank.get('file_sha256', ''))
    if not re.fullmatch('[0-9a-f]{64}', digest):
        raise ValueError('Missing source workbook SHA-256')
    return {'schema': SCHEMA, 'file_name': Path(str(bank.get('file_name') or 'SLS deflection.xlsx')).name,
        'file_sha256': digest, 'component': 'WebTop Vert',
        'source_contract': 'User-supplied SLS girder envelope; native Max/Min pairs',
        'case_name': None, 'load_definition_in_export': False,
        'ignored_sheets': list(bank.get('ignored_sheets') or []), 'girders': sheets}


def parse_workbook(payload, file_name='SLS deflection.xlsx'):
    """Read an upload atomically; no application state is mutated on failure."""
    if Path(file_name).suffix.casefold() != '.xlsx':
        raise ValueError('Upload the Envelope SLS Girder Deflection .xlsx workbook')
    try:
        workbook = load_workbook(BytesIO(payload), read_only=True, data_only=True)
    except (ValueError, OSError, BadZipFile, KeyError):
        raise ValueError('Cannot read this .xlsx workbook; export the SLS girder envelope again') from None
    sheets, ignored = [], []
    try:
        for ws in workbook:
            rows = iter(ws.iter_rows(values_only=True))
            header, header_row = None, None
            for row_number, row in enumerate(rows, 1):
                labels = [_text(cell).casefold() for cell in row]
                if all(name.casefold() in labels for name in REQUIRED):
                    if any(labels.count(name.casefold()) != 1 for name in REQUIRED):
                        raise ValueError(ws.title + ': duplicate required column header')
                    header = {name: labels.index(name.casefold()) for name in REQUIRED}
                    header_row = row_number
                    break
                if row_number >= 25:
                    break
            if header is None:
                ignored.append(ws.title)
                continue
            units = next(rows, ())
            get = lambda row, name: row[header[name]] if len(row) > header[name] else None
            sheet = {'girder': ws.title, 'distance_unit': get(units, 'Girder Distance'),
                'vertical_unit': get(units, 'WebTop Vert'), 'rows': []}
            for row_number, row in enumerate(rows, header_row + 2):
                if not any(value is not None and _text(value) != '' for value in row):
                    continue
                sheet['rows'].append({'source_row': row_number, **{name: get(row, name) for name in REQUIRED}})
            sheets.append(_validate_sheet(sheet))
    finally:
        workbook.close()
    if not sheets:
        raise ValueError('Only Envelope SLS Girder Deflection is supported. Required columns: '
            + ', '.join(REQUIRED) + '; include the units row and Max/Min at every station.')
    return validate_bank({'schema': SCHEMA, 'file_name': file_name,
        'file_sha256': sha256(payload).hexdigest(), 'girders': sheets, 'ignored_sheets': ignored})


def input_metadata(state):
    bank = state.get(BANK_KEY)
    return deepcopy(validate_bank(bank)) if bank is not None else None
