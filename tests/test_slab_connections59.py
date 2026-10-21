"""Geometry-driven exclusions cannot leak connection peaks into design."""
from copy import deepcopy
from io import BytesIO
from unittest.mock import patch

import numpy as np
import openpyxl
import pandas as pd
import pytest

from concrete_pmm_pro.analysis import slab_shell56 as engine
from concrete_pmm_pro.analysis.slab_connections59 import annotate, layout, INCLUDED, REASON, _merged
from concrete_pmm_pro.analysis.slab_shell56_geometry import recover
from concrete_pmm_pro.io.slab_shell_import import BANK_KEY, parse_files
from concrete_pmm_pro.visualization.slab_shell56 import element_envelope, plan_figure, VIEWS
from test_slab_shell55 import fixtures, upload
from test_slab_shell56 import model_state


def connected_state():
    state = model_state()
    state['section_preset_key'] = 'parametric_i_girder'
    state['section_parameters'].update(B1_mm=600., D1_mm=1500.)
    state[BANK_KEY].tables['frames'] = pd.DataFrame([
        dict(Frame='G1', JointI='~1', JointJ='A', IsCurved='No'),
        dict(Frame='G2', JointI='A', JointJ='B', IsCurved='No')])
    return state


def test_complete_existing_workbook_can_supply_optional_frame_table():
    book = openpyxl.load_workbook(BytesIO(upload(fixtures()).getvalue()))
    sheet = book.create_sheet('Connectivity - Frame')
    sheet.append(['TABLE:  Connectivity - Frame'])
    sheet.append(['Frame','JointI','JointJ','IsCurved'])
    sheet.append(['Text','Text','Text','Yes/No'])
    sheet.append(['G1','1','2','No'])
    data = BytesIO(); book.save(data); data.name = 'Complete CSI export.xlsx'
    package = parse_files([data])
    assert package.ready, package.errors
    assert package.tables['frames'].iloc[0].to_dict() == dict(Frame='G1',JointI='1',JointJ='2',IsCurved='No')


def test_mask_uses_lines_and_export_tolerance_not_a_physical_band_or_load_size():
    frame = pd.DataFrame({'X (m)':[.5,.5,2.01,0,.5], 'Y (m)':[0,5e-5,0,.5,-1e-5],
                          'M11':[1.,1e9,1.,1.,1.]})
    lines = pd.DataFrame([dict(Connection='Girder',ID='G',**{'X1 (m)':0,'Y1 (m)':0,'X2 (m)':2,'Y2 (m)':0}),
                          dict(Connection='Diaphragm',ID='D',**{'X1 (m)':0,'Y1 (m)':0,'X2 (m)':0,'Y2 (m)':1})])
    result = annotate(frame, {'lines':lines})
    assert result[INCLUDED].tolist() == [False,True,True,False,False]
    np.testing.assert_array_equal(result['M11'], frame['M11'])


def test_collinear_girder_segments_merge_without_excluding_the_gap():
    lines = _merged([('a',np.array([0.,0.]),np.array([1.,0.])),
                     ('b',np.array([1.,0.]),np.array([2.,0.])),
                     ('c',np.array([3.,0.]),np.array([4.,0.]))], 'Girder', np.array([1.,0.]))
    assert len(lines) == 2
    frame = pd.DataFrame({'X (m)':[1.5,2.5,3.5], 'Y (m)':[0.,0.,0.]})
    assert annotate(frame, {'lines':pd.DataFrame(lines)})[INCLUDED].tolist() == [False,True,False]


def test_excluded_records_are_not_solved_and_have_no_design_resistance_or_dc():
    state = connected_state(); source = state[BANK_KEY].tables['forces'].copy(deep=True)
    with patch.object(engine, 'section_check', wraps=engine.section_check) as solver:
        result = engine.calculate(state)
    assert sum(len(c.args[0]) for c in solver.call_args_list) == 32
    assert result['source_rows'] == 64 and result['design_rows'] == 32 and result['excluded_rows'] == 32
    excluded = result['rows'].loc[~result['rows'][INCLUDED]]
    assert excluded['Field D/C'].isna().all() and excluded['Concrete D/C'].isna().all()
    for group in engine.GROUPS:
        assert excluded[group+' As required (mm²/m)'].isna().all()
    assert excluded['Status'].eq('EXCLUDED').all()
    assert not result['incomplete']
    assert result['coverage']['Missing records'].sum() == 0
    assert result['counts'].get('EXCLUDED',0) == 0
    pd.testing.assert_frame_equal(state[BANK_KEY].tables['forces'], source)
    assert not pd.Series(result['governing'])[REASON]


def test_arbitrarily_large_excluded_actions_cannot_change_design_summaries_or_status():
    state = connected_state(); first = engine.calculate(state)
    other = deepcopy(state)
    forces = other[BANK_KEY].tables['forces']
    excluded_joints = set(first['rows'].loc[~first['rows'][INCLUDED],'Joint'])
    cols = ['F11','F22','F12','M11','M22','M12','V13','V23']
    mask = forces['Joint'].isin(excluded_joints)
    forces.loc[mask,cols] = 1e12
    second = engine.calculate(other)
    assert (second['rows'].loc[~second['rows'][INCLUDED],cols] == 1e12).all().all()
    pd.testing.assert_frame_equal(first['summary'], second['summary'])
    pd.testing.assert_frame_equal(first['check_summary'], second['check_summary'])
    assert first['governing'] == second['governing'] or first['governing']['Field D/C'] == second['governing']['Field D/C']
    assert first['status'] == second['status']


def test_map_uses_retained_corners_and_keeps_complete_import_coverage():
    result = engine.calculate(connected_state())
    for view in VIEWS:
        selected = element_envelope(result,view=view)
        assert selected[INCLUDED].all()
        assert not selected['Map incomplete'].any()
        assert selected['Excluded connection records'].eq(16).all()
        fig = plan_figure(result,view=view)
        assert any(t.name=='I-girder CL' for t in fig.data)
        assert all(not t.name.startswith('Gov.') or all(y != 0 for y in t.y) for t in fig.data)


def test_all_excluded_element_is_grey_without_a_zero_dc_or_false_missing_import():
    state = connected_state()
    # Add only the first element's opposite edge. Its four corners are then
    # excluded, while the second element still has retained corner records.
    state[BANK_KEY].tables['frames'] = pd.concat([state[BANK_KEY].tables['frames'],
        pd.DataFrame([dict(Frame='G3',JointI='4',JointJ='D',IsCurved='No')])], ignore_index=True)
    result = engine.calculate(state)
    selected = element_envelope(result,view='N–M field')
    assert selected.loc['1-1','Map excluded']
    assert pd.isna(selected.loc['1-1','Map D/C'])
    assert not selected.loc['1-1','Map incomplete']
    fig = plan_figure(result,view='N–M field')
    assert any('No nominal slab D/C' in str(t.text) for t in fig.data)


def test_no_remaining_nodes_cannot_produce_pass_or_zero_strength():
    state = connected_state()
    state[BANK_KEY].tables['frames'] = pd.concat([state[BANK_KEY].tables['frames'],
        pd.DataFrame([dict(Frame='G3',JointI='4',JointJ='C',IsCurved='No')])], ignore_index=True)
    with pytest.raises(ValueError, match='No nominal slab nodes'):
        engine.calculate(state)


def test_flange_edge_and_negative_moment_reference_are_distinct_and_invalidate_stale_location():
    state = connected_state(); geometry = recover(state[BANK_KEY])
    support = layout(state[BANK_KEY],geometry,state)
    assert support['top_flange_mm'] == 600 and support['negative_moment_offset_mm'] == 200
    assert support['top_flange_mm']/2 == 300
    state[engine.RESULT_KEY] = engine.calculate(state)
    state['section_parameters']['B1_mm'] = 900
    assert engine.current(state) is None


def test_rotated_girder_line_and_transverse_diaphragm_mask_in_plan_coordinates():
    angle = np.deg2rad(31.)
    rotation = np.array([[np.cos(angle),-np.sin(angle)],[np.sin(angle),np.cos(angle)]])
    p = np.array([[.5,0],[.5,.1],[0,.5]]) @ rotation.T
    a,b = np.array([[0.,0.],[2.,0.]]) @ rotation.T
    c,d = np.array([[0.,0.],[0.,1.]]) @ rotation.T
    lines = pd.DataFrame([dict(Connection=k,ID=k,**dict(zip(['X1 (m)','Y1 (m)','X2 (m)','Y2 (m)'],[*u,*v])))
                           for k,u,v in [('Girder',a,b),('Diaphragm',c,d)]])
    assert annotate(pd.DataFrame(p,columns=['X (m)','Y (m)']),{'lines':lines})[INCLUDED].tolist() == [False,True,False]
