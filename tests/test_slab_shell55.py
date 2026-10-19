"""Shell input contract: complete native tables, correspondence and UI routing."""
from copy import deepcopy
from io import BytesIO
from pathlib import Path
from unittest.mock import patch

import openpyxl
import pytest
from streamlit.testing.v1 import AppTest

from concrete_pmm_pro.core.analysis import AnalysisModeSettings
from concrete_pmm_pro.core.project import ProjectModel
from concrete_pmm_pro.io import slab_shell_import as shell
from concrete_pmm_pro.io.project_io import apply_project_to_session_state
from concrete_pmm_pro.ui.navigation import analysis_subpages_for_session, loads_subpages_for_session
from concrete_pmm_pro.ui.slab_reinforcement import export_guide, reinforcement_inventory

ROOT = Path(__file__).resolve().parents[1]


def fixtures():
    values = {
        "forces": [dict(Area=1, AreaElem="1-1", ShellType="Shell-Thin", Joint="~1", OutputCase=case,
            CaseType="Combination", StepType=step, **dict(zip(shell.FORCES, range(1, 9))))
            for case in ("U1A", "U1B", "U2A", "U2B") for step in ("Max", "Min")],
        "elements": [dict(AreaElem="1-1", AreaObject=1, ElemJt1="~1", ElemJt2="2", ElemJt3="3", ElemJt4="4")],
        "nodes": [dict(JointElem=j, GlobalX=x, GlobalY=y, GlobalZ=0) for j, x, y in (("~1", 0, 0), (2, 1, 0), (3, 1, 1), (4, 0, 1))],
        "flags": [dict(Area=1, BridgeObj="BOBJ1", BOSpan="Span1", CompType="Top Slab", StartSta=0., EndSta=1.)],
        "thickness": [dict(Area=1, ThickOver="Object", Thick1=.2, Thick2=.35, Thick3=.35, Thick4=.2)],
        "sections": [dict(Area=1, Section="ASEC1", MatProp="Default")],
        "axes": [dict(Area=1, Angle=0, AdvanceAxes="Yes")],
        "advanced": [dict(Area=1, LocalPlane="31", PlOption1="Two Joints", PlCoordSys="GLOBAL",
            CoordDir1="Z", CoordDir2="Y", PlVecJt1="~1", PlVecJt2=2, PlVecX=1, PlVecY=0, PlVecZ=0)],
        "parents": [dict(Area=1, NumJoints=4, Joint1="~1", Joint2=2, Joint3=3, Joint4=4)],
        "joints": [dict(Joint=j, GlobalX=x, GlobalY=y, GlobalZ=0) for j, x, y in ((2, 1, 0), (3, 1, 1), (4, 0, 1))],
    }
    return values


def upload(data=None, *, omitted=(), units_override=None, name="renamed complete model.xlsx", extras=True):
    data = deepcopy(data or fixtures())
    book = openpyxl.Workbook()
    book.remove(book.active)
    for spec in shell.SPECS:
        if spec.key in omitted:
            continue
        # Deliberately rename every worksheet; detection must use its schema.
        sheet = book.create_sheet(spec.key)
        sheet.append(["TABLE:  " + spec.title])
        columns = list(spec.columns)
        if spec.key == "forces":
            columns += ["MMax", "SillyUnusedValue"]
        sheet.append(columns)
        units = {c: "Text" for c in columns}
        for c in shell.FORCES:
            units[c] = "KN-m/m" if c.startswith("M") else "KN/m"
        for c in ("GlobalX", "GlobalY", "GlobalZ", "Thick1", "Thick2", "Thick3", "Thick4", "StartSta", "EndSta"):
            units[c] = "m"
        units["Angle"] = "Degrees"
        if units_override and spec.key in units_override:
            units.update(units_override[spec.key])
        sheet.append([units[c] for c in columns])
        for row in data[spec.key]:
            sheet.append([row.get(c, 0 if c == "MMax" else None) for c in columns])
    if extras:
        sheet = book.create_sheet("Connectivity - Frame")
        sheet.append(["TABLE:  Connectivity - Frame"])
        sheet.append(["Frame", "JointI", "JointJ"])
        sheet.append(["Text", "Text", "Text"])
        sheet.append(["F1", "2", "3"])
    result = BytesIO()
    book.save(result)
    result.name = name
    return result


def test_complete_renamed_workbook_preserves_coupled_rows_auxiliaries_and_thickness():
    result = shell.parse_files([upload()])
    assert result.ready and not result.errors and not result.missing
    assert result.summary["cases"] == ["U1A", "U1B", "U2A", "U2B"]
    assert result.summary["slab_force_rows"] == 8
    assert result.summary["thickness_min_mm"] == 200
    assert result.summary["thickness_max_mm"] == 350
    assert len(result.tables["forces"]) == 8
    assert tuple(result.tables["forces"].iloc[0][list(shell.FORCES)]) == tuple(range(1, 9))
    assert result.tables["forces"]["MMax"].eq(0).all()
    assert any(s["Use"] == "Not used" for s in result.sources)
    assert result.tables["nodes"]["JointElem"].iloc[0] == "~1"


def test_advanced_axes_required_only_when_active():
    result = shell.parse_files([upload(omitted=("advanced",))])
    assert not result.ready and [s.key for s in result.missing] == ["advanced"]
    data = fixtures()
    data["axes"][0]["AdvanceAxes"] = "No"
    result = shell.parse_files([upload(data, omitted=("advanced",))])
    assert result.ready


def test_generated_joint_coordinates_cannot_be_replaced_with_original_joints_only():
    data = fixtures()
    data["nodes"] = data["nodes"][1:]
    result = shell.parse_files([upload(data)])
    assert not result.ready
    assert any("analysis mesh joints" in e for e in result.errors)


@pytest.mark.parametrize("failure", ["component", "parent", "corner", "duplicate", "nonfinite", "thickness", "reference"])
def test_input_failures_are_blocked(failure):
    data = fixtures()
    if failure == "component": data["forces"][0]["M12"] = None
    if failure == "parent": data["forces"][0]["Area"] = "99"
    if failure == "corner": data["forces"][0]["Joint"] = "other"
    if failure == "duplicate": data["forces"].append(deepcopy(data["forces"][0]))
    if failure == "nonfinite": data["nodes"][0]["GlobalX"] = "inf"
    if failure == "thickness": data["thickness"][0]["Thick1"] = 0
    if failure == "reference": data["advanced"][0]["PlVecJt2"] = "not-in-export"
    result = shell.parse_files([upload(data)])
    assert not result.ready and result.errors


def test_incompatible_units_rejected_instead_of_silently_rescaled():
    result = shell.parse_files([upload(units_override={"forces": {"M11": "N-mm/mm"}})])
    assert not result.ready and any("unsupported unit for M11" in e for e in result.errors)


def test_identical_files_and_tables_are_deduplicated_conflicts_blocked():
    first = upload()
    result = shell.parse_files([first, first, upload(name="same tables different container.xlsx", extras=False)])
    assert result.ready and len(result.tables["forces"]) == 8 and result.notes
    different = fixtures()
    different["forces"][0]["F11"] = 999
    result = shell.parse_files([first, upload(different)])
    assert not result.ready and any("Conflicting copies" in e for e in result.errors)


def test_diaphragm_records_stay_in_bank_but_not_in_slab_selection():
    data = fixtures()
    data["flags"].append({**data["flags"][0], "Area": 2, "CompType": "Diaphragm"})
    data["elements"].append({**data["elements"][0], "AreaElem": "2-1", "AreaObject": 2})
    data["forces"].append({**data["forces"][0], "Area": 2, "AreaElem": "2-1"})
    result = shell.parse_files([upload(data)])
    assert result.ready
    assert len(result.tables["forces"]) == 9 and result.summary["slab_force_rows"] == 8
    assert result.summary["excluded_force_rows"] == 1


def test_stale_file_selection_is_not_exposed_as_current_data():
    result = shell.parse_files([upload()])
    state = {shell.BANK_KEY: result, shell.PENDING_KEY: result.signature}
    assert shell.current_package(state) is result
    state[shell.PENDING_KEY] = ("changed",)
    assert shell.current_package(state) is None


def test_new_project_clears_session_shell_package_and_does_not_persist_in_json():
    state = {shell.BANK_KEY: "old model", shell.PENDING_KEY: ("old",), "slab_shell55_uploads_0": "old widget"}
    apply_project_to_session_state(ProjectModel(), state)
    assert not any(str(k).startswith("slab_shell55_") for k in state)


def test_navigation_is_bridge_scoped_and_existing_workflows_remain_available():
    for member in ("beam_girder", "building_beam_girder", "portal_frame_crossbeam", "column_pier_pmm"):
        state = {"analysis_mode_settings": AnalysisModeSettings(member_type=member)}
        assert ("Slab Reinforcement" in analysis_subpages_for_session(state)) == (member == "beam_girder")
        assert ("Slab Shell Results" in loads_subpages_for_session(state)) == (member == "beam_girder")


def test_export_instructions_name_supplied_files_and_required_tables():
    guide = export_guide()
    for row in shell.FILE_GUIDE:
        assert row["File / export set"] in guide
    for spec in shell.SPECS:
        assert spec.title in guide or spec.key == "advanced"
    assert "moving-load recording-settings" in guide and "No stress workbook" in guide


def test_physical_slab_bars_are_read_even_if_girder_flexure_inclusion_is_off():
    params = {"Tslab_mm": 350, "deck_slab_layer_order": "transverse_outer_longitudinal_inner",
        "deck_long_rebar_fy_MPa": 390, "include_deck_long_rebar_in_positive_Mn": False}
    for face in ("top", "bottom"):
        params.update({f"deck_long_rebar_{face}_diameter_mm": 16, f"deck_long_rebar_{face}_spacing_mm": 150})
    cfg = {"slab_layer_order": "transverse_outer_longitudinal_inner", "slab_transverse_mesh": {
        face: {"Diameter_mm": 16, "Spacing_mm": 150, "Clear_cover_mm": 40, "fy_MPa": 390} for face in ("top", "bottom")}}
    state = {"section_parameters": params, "beam_girder_interface_shear_settings": {"dowel_connection": cfg}}
    frame = reinforcement_inventory(state)
    assert len(frame) == 4 and frame["Input status"].eq("Available").all()
    assert frame.loc[frame["Direction"] == "Longitudinal", "Clear cover (mm)"].eq(56).all()


def test_actual_app_slab_routes_do_not_run_girder_solvers_or_reparse_on_navigation():
    state = {}
    apply_project_to_session_state(ProjectModel(analysis_mode_settings=AnalysisModeSettings(member_type="beam_girder")), state)
    state[shell.BANK_KEY] = shell.parse_files([upload()])
    app = AppTest.from_file(str(ROOT / "app.py"), default_timeout=60)
    for key, value in state.items():
        app.session_state[key] = deepcopy(value)
    app.session_state["_nav_active_workspace"] = "Loads"
    app.session_state["_nav_loads_subpage"] = "Slab Shell Results"
    with patch.object(shell, "parse_files", side_effect=AssertionError("unexpected parse")):
        app.run()
        assert not app.exception
        assert any("Import & validate" in b.label for b in app.button)
        next(b for b in app.button if b.key == "slab_shell55_open_analysis").click().run()
        assert not app.exception
        assert app.session_state["_nav_analysis_subpage"] == "Slab Reinforcement"
        assert any("NOT CALCULATED" in i.value for i in app.info)
        assert next(b for b in app.button if b.key=='slab_shell55_strength56_calculate').disabled
        next(b for b in app.button if b.key == "slab_shell55_to_loads").click().run()
        assert not app.exception
        assert app.session_state["_nav_loads_subpage"] == "Slab Shell Results"
