"""Native CSI static exports may omit StepType; engineering inputs stay intact."""
from copy import deepcopy
from io import BytesIO

import numpy as np
import openpyxl
import pandas as pd
import pytest

from concrete_pmm_pro.io import slab_shell_import as source
from test_slab_shell55 import fixtures, upload


def static_data():
    data = fixtures()
    data['forces'] = [r for r in data['forces'] if r['StepType'] == 'Max']
    for row in data['forces']:
        row['StepType'] = None
    return data


def native_upload(data=None, *, remove=('StepType',), units_override=None, step_numbers=None):
    original = upload(data or static_data(), extras=False, units_override=units_override)
    book = openpyxl.load_workbook(original)
    sheet = book['forces']
    if step_numbers is not None:
        column = sheet.max_column + 1
        sheet.cell(2, column, 'StepNum')
        sheet.cell(3, column, 'Unitless')
        for row, value in enumerate(step_numbers, 4):
            sheet.cell(row, column, value)
    for header in remove:
        columns = [c.value for c in sheet[2]]
        sheet.delete_cols(columns.index(header) + 1)
    result = BytesIO()
    book.save(result)
    book.close()
    result.name = 'native static export.xlsx'
    return result


def test_omitted_step_header_imports_all_static_actions_without_inventing_envelopes():
    data = static_data()
    result = source.parse_files([native_upload(data)])
    assert result.ready and not result.errors and not result.missing
    forces = result.tables['forces']
    assert len(forces) == 4 and forces.StepType.eq('').all()
    assert forces.CaseType.eq('Combination').all()
    assert list(forces.OutputCase) == [r['OutputCase'] for r in data['forces']]
    np.testing.assert_array_equal(forces[list(source.FORCES)],
        np.array([[r[c] for c in source.FORCES] for r in data['forces']]))
    assert 'MMax' in forces and 'StepNum' not in forces
    assert result.units['forces']['StepType'] == 'Text'
    assert any('StepType not exported' in n for n in result.notes)
    assert result.summary['case_rows'].StepType.eq('').all()


def test_explicit_blank_and_omitted_step_tables_are_equivalent_duplicates():
    result = source.parse_files([native_upload(remove=()), native_upload()])
    assert result.ready
    assert len(result.tables['forces']) == 4
    assert any('Identical duplicate table ignored' in n for n in result.notes)


@pytest.mark.parametrize('missing', ['F12', 'M12', 'V23', 'OutputCase', 'CaseType'])
def test_real_required_fields_still_block_with_precise_missing_column_error(missing):
    result = source.parse_files([native_upload(remove=('StepType', missing))])
    assert not result.ready and 'forces' not in result.tables
    assert any(e.endswith('missing required columns: ' + missing) for e in result.errors)
    assert all('StepType' not in e for e in result.errors)


def test_ambiguous_repeated_static_corners_are_blocked_without_merging():
    data = static_data()
    repeated = deepcopy(data['forces'][0])
    repeated['M12'] = 999.
    data['forces'].append(repeated)
    result = source.parse_files([native_upload(data)])
    assert not result.ready
    assert any('duplicate area / element / corner / case / step' in e for e in result.errors)
    assert len(result.tables['forces']) == 5
    assert result.tables['forces'].M12.iloc[-1] == 999.


def test_exported_step_numbers_are_retained_and_distinguish_rows():
    data = static_data()
    data['forces'] = [deepcopy(data['forces'][0]), deepcopy(data['forces'][0])]
    data['forces'][1]['V13'] = 42.
    result = source.parse_files([native_upload(data, step_numbers=[0., 1.])])
    assert result.ready and result.tables['forces'].StepType.eq('').all()
    assert result.tables['forces'].StepNum.tolist() == [0., 1.]
    assert result.tables['forces'].V13.tolist() == [7., 42.]


def test_static_import_still_rejects_incompatible_units():
    result = source.parse_files([native_upload(units_override={'forces': {'V13': 'N/mm'}})])
    assert not result.ready
    assert any('unsupported unit for V13' in e for e in result.errors)


def test_max_min_export_preserves_step_names_and_complete_concurrent_vectors():
    data = fixtures()
    result = source.parse_files([upload(data, extras=False)])
    assert result.ready and set(result.tables['forces'].StepType) == {'Max', 'Min'}
    expected = pd.DataFrame(data['forces'])
    for column in (*source.FORCES, 'OutputCase', 'CaseType', 'StepType'):
        np.testing.assert_array_equal(result.tables['forces'][column], expected[column])
    assert not any('StepType not exported' in n for n in result.notes)
